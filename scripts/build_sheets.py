"""Build one blank count workbook per rater.

Usage (from the G15 repo root):
    python scripts/build_sheets.py <path-to-EMSE2025SAR-Replication>

Creates count_sheets/rater1_counts.xlsx and count_sheets/rater2_counts.xlsx.
Each has one tab per app with two tables:
  A. every ground-truth item  -> rater writes TP or FN
  B. every Code2DFD output item -> rater writes MATCH or FP
The sheets contain NO labels from the paper.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from common import APPS, ground_truth, tool_output

HEAD = Font(bold=True, color="FFFFFF")
FILL_A = PatternFill("solid", fgColor="1F4E78")
FILL_B = PatternFill("solid", fgColor="7F3F00")
INPUT = PatternFill("solid", fgColor="FFF2CC")


def build(repo, rater, out):
    wb = Workbook()
    info = wb.active
    info.title = "README"
    lines = [
        f"CS4928 G15 / P21 - count sheet for {rater.upper()}",
        "",
        "1. Read docs/counting_rule.md before starting.",
        "2. Open the raw JSON and the ground-truth .txt named on each tab. Nothing else.",
        "3. Do NOT open processed_results/ or merged_results/ until both raters are finished.",
        "4. Do NOT talk to the other rater about any item until both files are committed.",
        "5. Fill only the yellow cells. Use the drop-downs.",
        "   Table A (ground truth): TP = Code2DFD reported it, FN = it did not.",
        "   Table B (tool output):  MATCH = it is a ground-truth item, FP = it is not.",
        "6. Put a short reason in Note for every judgement call.",
        "7. Save, then commit this file unchanged. Never edit it after committing.",
    ]
    for i, t in enumerate(lines, 1):
        info.cell(i, 1, t)
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 100
    info.cell(len(lines) + 2, 1, "Rater name:").font = Font(bold=True)
    info.cell(len(lines) + 2, 2).fill = INPUT
    info.cell(len(lines) + 3, 1, "Date finished:").font = Font(bold=True)
    info.cell(len(lines) + 3, 2).fill = INPUT

    for app in APPS:
        ws = wb.create_sheet(app[:31])
        gt, tool = ground_truth(repo, app), tool_output(repo, app)
        ws["A1"] = f"App: {app}"
        ws["A1"].font = Font(bold=True, size=12)
        ws["A2"] = f"Inputs: groundtruth_textual/{app}.txt  +  raw_results/Code2DFD/{app}-code2DFD-raw.json"

        dv_a = DataValidation(type="list", formula1='"TP,FN"', allow_blank=True)
        dv_b = DataValidation(type="list", formula1='"MATCH,FP"', allow_blank=True)
        ws.add_data_validation(dv_a); ws.add_data_validation(dv_b)

        r = 4
        ws.cell(r, 1, "A. GROUND-TRUTH ITEMS  (label TP or FN)").font = Font(bold=True)
        r += 1
        for c, h in enumerate(["ID", "Characteristic", "Ground-truth item", "Label", "Matching tool item (if TP)", "Note"], 1):
            cell = ws.cell(r, c, h); cell.font = HEAD; cell.fill = FILL_A
        for i, (ch, item) in enumerate(gt, 1):
            r += 1
            ws.cell(r, 1, f"G{i:03d}"); ws.cell(r, 2, ch); ws.cell(r, 3, item)
            for c in (4, 5, 6):
                ws.cell(r, c).fill = INPUT
            dv_a.add(ws.cell(r, 4))

        r += 2
        ws.cell(r, 1, "B. CODE2DFD OUTPUT ITEMS  (label MATCH or FP)").font = Font(bold=True)
        r += 1
        for c, h in enumerate(["ID", "Characteristic", "Tool output item", "Label", "Matching ground-truth ID (if MATCH)", "Note"], 1):
            cell = ws.cell(r, c, h); cell.font = HEAD; cell.fill = FILL_B
        for i, (ch, item) in enumerate(tool, 1):
            r += 1
            ws.cell(r, 1, f"T{i:03d}"); ws.cell(r, 2, ch); ws.cell(r, 3, item)
            for c in (4, 5, 6):
                ws.cell(r, c).fill = INPUT
            dv_b.add(ws.cell(r, 4))

        for col, w in zip("ABCDEF", (7, 14, 55, 9, 38, 50)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A4"
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (3, 5, 6))
    wb.save(out)
    return {app: (len(ground_truth(repo, app)), len(tool_output(repo, app))) for app in APPS}


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    repo = sys.argv[1]
    Path("count_sheets").mkdir(exist_ok=True)
    for n in (1, 2):
        sizes = build(repo, f"rater{n}", Path("count_sheets") / f"rater{n}_counts.xlsx")
    print("Built count_sheets/rater1_counts.xlsx and rater2_counts.xlsx")
    for app, (g, t) in sizes.items():
        print(f"  {app:34s} ground-truth items: {g:3d}   tool items: {t:3d}")
