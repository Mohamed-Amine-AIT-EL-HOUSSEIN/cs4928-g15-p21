"""Agreement, scoring and comparison for the G15 M2 reproduction.

Step 1 - after BOTH raters committed their sheets:
    python scripts/score.py agreement
      -> results/agreement.md, results/conflicts.csv
      -> count_sheets/agreed_counts.xlsx  (agreed labels filled in, conflicts left
         blank and highlighted red for the group to resolve)

Step 2 - after every conflict in agreed_counts.xlsx is resolved:
    python scripts/score.py score <path-to-EMSE2025SAR-Replication>
      -> results/m2_results.csv, results/m2_summary.md
      (this is the first step that reads the paper's own labels)
"""
import csv
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from common import APPS, CHARS, paper_counts, prf

SHEETS = Path("count_sheets")
RESULTS = Path("results")
RED = PatternFill("solid", fgColor="F4B6B6")

# Frozen M1 decision rule (workbook G15, column H) + clarification D-02
POOLED_TOL = 0.05     # pooled P, R, F1 within +/-0.05 of the paper
APP_F1_TOL = 0.10     # every app's F1 within +/-0.10 of the paper
KAPPA_MIN = 0.60      # lower of kappa(A) and kappa(B) must be >= 0.60


def read_sheet(path):
    """{(app, id): {'char', 'item', 'label', 'match', 'note', 'row'}}"""
    wb = load_workbook(path)
    out = {}
    for app in APPS:
        ws = wb[app[:31]]
        for row in ws.iter_rows(min_row=1):
            rid = row[0].value
            if isinstance(rid, str) and rid[:1] in "GT" and rid[1:].isdigit():
                lab = (row[3].value or "").strip().upper()
                out[(app, rid)] = {"char": row[1].value, "item": row[2].value,
                                   "label": lab, "match": row[4].value,
                                   "note": row[5].value, "row": row[0].row}
    return out


def cohen_kappa(pairs):
    n = len(pairs)
    if n == 0:
        return None, None
    po = sum(a == b for a, b in pairs) / n
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / (n * n)
    kappa = None if pe == 1 else (po - pe) / (1 - pe)
    return po, kappa


def fmt(x, nd=2):
    return "n/a" if x is None else f"{x:.{nd}f}"


def agreement():
    r1 = read_sheet(SHEETS / "rater1_counts.xlsx")
    r2 = read_sheet(SHEETS / "rater2_counts.xlsx")
    blanks = [k for k in r1 if not r1[k]["label"] or not r2[k]["label"]]
    if blanks:
        sys.exit(f"{len(blanks)} unlabelled rows (e.g. {blanks[:5]}). Both raters must finish first.")

    RESULTS.mkdir(exist_ok=True)
    lines = ["# Inter-rater agreement (Cohen's kappa)", "",
             "Table A = ground-truth items (TP/FN). Table B = Code2DFD output items (MATCH/FP).", "",
             "| Scope | Table | Items | % agree | kappa |", "|---|---|---|---|---|"]
    conflicts = []
    for scope in APPS + ["ALL 4 APPS"]:
        for table in ("A", "B"):
            pairs = [(r1[k]["label"], r2[k]["label"]) for k in r1
                     if k[1][0] == ("G" if table == "A" else "T")
                     and (scope == "ALL 4 APPS" or k[0] == scope)]
            po, ka = cohen_kappa(pairs)
            lines.append(f"| {scope} | {table} | {len(pairs)} | {fmt(po and po * 100, 1)} | {fmt(ka)} |")
    pooled = {t: cohen_kappa([(r1[k]["label"], r2[k]["label"]) for k in r1
                              if k[1][0] == ("G" if t == "A" else "T")])[1] for t in "AB"}
    decision_kappa = min(v for v in pooled.values() if v is not None) if any(
        v is not None for v in pooled.values()) else None
    lines += ["", f"**Kappa used by the decision rule (lower of A and B, all apps): {fmt(decision_kappa)}**",
              f"(threshold >= {KAPPA_MIN}; expected in M1: >= 0.8)", "",
              "Note: kappa is n/a when both raters gave every item the same label "
              "(no variation), so % agreement is reported alongside it."]

    # agreed workbook = rater1's file with conflicts blanked + highlighted
    wb = load_workbook(SHEETS / "rater1_counts.xlsx")
    wb["README"]["A1"] = "CS4928 G15 / P21 - AGREED counts (resolve every red row, then run score.py score)"
    for k in r1:
        if r1[k]["label"] != r2[k]["label"]:
            app, rid = k
            conflicts.append([app, rid, r1[k]["char"], r1[k]["item"], r1[k]["label"],
                              r2[k]["label"], r1[k]["note"] or "", r2[k]["note"] or ""])
            ws = wb[app[:31]]
            row = r1[k]["row"]
            ws.cell(row, 4).value = None
            ws.cell(row, 6).value = (f"CONFLICT R1={r1[k]['label']} ({r1[k]['note'] or '-'}) | "
                                     f"R2={r2[k]['label']} ({r2[k]['note'] or '-'}) | RESOLUTION: ")
            for c in range(1, 7):
                ws.cell(row, c).fill = RED
    wb.save(SHEETS / "agreed_counts.xlsx")

    with open(RESULTS / "conflicts.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["app", "id", "characteristic", "item", "rater1", "rater2", "note1", "note2"])
        w.writerows(conflicts)
    lines += ["", f"Conflicts to resolve: {len(conflicts)} (see results/conflicts.csv)"]
    (RESULTS / "agreement.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines))


def score(repo):
    agreed = read_sheet(SHEETS / "agreed_counts.xlsx")
    blanks = [k for k, v in agreed.items() if not v["label"]]
    if blanks:
        sys.exit(f"{len(blanks)} conflicts still unresolved in agreed_counts.xlsx (e.g. {blanks[:5]}).")

    ours = {a: {c: {"TP": 0, "FP": 0, "FN": 0} for c in CHARS} for a in APPS}
    for (app, rid), v in agreed.items():
        lab = v["label"]
        if rid[0] == "G" and lab in ("TP", "FN"):
            ours[app][v["char"]][lab] += 1
        elif rid[0] == "T" and lab == "FP":
            ours[app][v["char"]]["FP"] += 1
    paper = {a: paper_counts(repo, a) for a in APPS}

    def total(src, apps, chars):
        t = Counter()
        for a in apps:
            for c in chars:
                t.update(src[a][c])
        return t["TP"], t["FP"], t["FN"]

    rows = []
    for scope in APPS + ["POOLED"]:
        apps = APPS if scope == "POOLED" else [scope]
        for ch in CHARS + ["all"]:
            chars = CHARS if ch == "all" else [ch]
            o, p = total(ours, apps, chars), total(paper, apps, chars)
            op, pp = prf(*o), prf(*p)
            rows.append([scope, ch, *o, *[fmt(x, 3) for x in op], *p, *[fmt(x, 3) for x in pp],
                         *[fmt(a - b if a is not None and b is not None else None, 3)
                           for a, b in zip(op, pp)]])
    RESULTS.mkdir(exist_ok=True)
    head = ["scope", "characteristic", "our_TP", "our_FP", "our_FN", "our_P", "our_R", "our_F1",
            "paper_TP", "paper_FP", "paper_FN", "paper_P", "paper_R", "paper_F1",
            "diff_P", "diff_R", "diff_F1"]
    with open(RESULTS / "m2_results.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(head); w.writerows(rows)

    # decision rule
    get = {(r[0], r[1]): r for r in rows}
    pooled = get[("POOLED", "all")]
    pooled_diffs = [float(x) for x in pooled[14:17]]
    app_f1_diffs = {a: float(get[(a, "all")][16]) for a in APPS}
    kappa_line = Path(RESULTS / "agreement.md").read_text(encoding="utf-8") \
        if (RESULTS / "agreement.md").exists() else ""
    kappa = None
    for line in kappa_line.splitlines():
        if line.startswith("**Kappa used by the decision rule"):
            val = line.split(":")[-1].strip(" *")
            kappa = None if val == "n/a" else float(val)

    checks = [
        ("Pooled P/R/F1 within +/-0.05", all(abs(d) <= POOLED_TOL for d in pooled_diffs),
         f"diffs P/R/F1 = {pooled_diffs}"),
        ("Every app F1 within +/-0.10", all(abs(d) <= APP_F1_TOL for d in app_f1_diffs.values()),
         f"{app_f1_diffs}"),
        ("Kappa >= 0.60", kappa is not None and kappa >= KAPPA_MIN, f"kappa = {fmt(kappa)}"),
    ]
    status = "REPRODUCED" if all(ok for _, ok, _ in checks) else "NOT REPRODUCED"
    md = ["# M2 result - Code2DFD on 4 benchmark apps", "",
          "| Scope | Our P | Our R | Our F1 | Paper P | Paper R | Paper F1 | diff F1 |",
          "|---|---|---|---|---|---|---|---|"]
    for r in rows:
        if r[1] == "all":
            md.append(f"| {r[0]} | {r[5]} | {r[6]} | {r[7]} | {r[11]} | {r[12]} | {r[13]} | {r[16]} |")
    md += ["", "## Decision rule (frozen in M1)", ""]
    md += [f"- {'PASS' if ok else 'FAIL'}: {name} ({detail})" for name, ok, detail in checks]
    md += ["", f"## Status: **{status}**", "",
           "Per-characteristic numbers: results/m2_results.csv"]
    (RESULTS / "m2_summary.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    print("\n".join(md))


if __name__ == "__main__":
    if len(sys.argv) >= 2 and sys.argv[1] == "agreement":
        agreement()
    elif len(sys.argv) == 3 and sys.argv[1] == "score":
        score(sys.argv[2])
    else:
        sys.exit(__doc__)
